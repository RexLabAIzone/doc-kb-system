import os
import struct
import subprocess
import time
import logging
import threading

import chardet

from config import setup_logging, MAX_FILE_SIZE_MB, CHECK_ENCODING, CHECK_OPENABLE, CHECK_CORRUPTION, SUSPICIOUS_CHARS_THRESHOLD

logger = setup_logging("checker")
_lock = threading.Lock()

MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024


def check_file(file_info):
    path = file_info["path"]
    category = file_info["category"]
    start = time.time()

    result = {
        "path": path,
        "name": file_info["name"],
        "ext": file_info["ext"],
        "size": file_info["size"],
        "category": category,
        "status": "ok",
        "issues": [],
        "encoding": None,
        "charset_confidence": None,
        "page_count": None,
        "image_dimensions": None,
        "duration_ms": 0.0,
        "ai_summary": None,
    }

    if file_info["size"] > MAX_FILE_SIZE_BYTES:
        result["status"] = "warning"
        result["issues"].append({
            "severity": "warning",
            "type": "file_too_large",
            "message": f"File exceeds size limit of {MAX_FILE_SIZE_MB}MB ({file_info['size'] / 1024 / 1024:.1f}MB)"
        })

    try:
        if category == "text":
            _check_text(path, result)
        elif category == "pdf":
            _check_pdf(path, result)
        elif category == "image":
            _check_image(path, result)
        elif category in ("office_doc", "office_spreadsheet", "office_presentation"):
            _check_office(path, result, category)
        elif category in ("audio", "video"):
            _check_audio_video(path, result)
        elif category == "archive":
            _check_archive(path, result)
        elif category == "ebook":
            _check_ebook(path, result)
        elif category in ("binary", "font"):
            _check_binary(path, result)
        else:
            _check_binary(path, result)
    except MemoryError:
        result["status"] = "error"
        result["issues"].append({
            "severity": "error",
            "type": "out_of_memory",
            "message": "Out of memory while checking file"
        })
    except Exception as e:
        result["status"] = "error"
        result["issues"].append({
            "severity": "error",
            "type": "check_error",
            "message": f"Unexpected error during check: {str(e)}"
        })

    if result["issues"]:
        has_error = any(i["severity"] == "error" for i in result["issues"])
        result["status"] = "error" if has_error else "warning"

    result["duration_ms"] = (time.time() - start) * 1000
    return result


def _read_file_preview(path, max_bytes=8192):
    try:
        with open(path, "rb") as f:
            return f.read(max_bytes)
    except Exception:
        return b""


def _read_full_content(path, max_bytes=10 * 1024 * 1024):
    try:
        size = os.path.getsize(path)
        if size > max_bytes:
            with open(path, "rb") as f:
                return f.read(max_bytes)
        with open(path, "rb") as f:
            return f.read()
    except Exception:
        return b""


def _detect_encoding(data):
    try:
        result = chardet.detect(data)
        return result["encoding"], result["confidence"]
    except Exception:
        return None, None


def _check_text(path, result):
    if CHECK_ENCODING:
        preview = _read_file_preview(path, 65536)
        encoding, confidence = _detect_encoding(preview)
        result["encoding"] = encoding
        result["charset_confidence"] = confidence
        if encoding is None or confidence is None:
            result["issues"].append({
                "severity": "warning",
                "type": "encoding_detection_failed",
                "message": "Could not detect text encoding"
            })
        elif confidence < 0.5:
            result["issues"].append({
                "severity": "warning",
                "type": "low_encoding_confidence",
                "message": f"Low encoding confidence: {confidence:.2f} ({encoding})"
            })
        elif encoding and encoding.lower() not in ("utf-8", "utf-16le", "ascii", "latin-1", "iso-8859-1", "windows-1252", "shift_jis", "euc-jp", "euc-kr", "gb2312", "gbk", "big5"):
            result["issues"].append({
                "severity": "info",
                "type": "uncommon_encoding",
                "message": f"Uncommon encoding detected: {encoding}"
            })

    if CHECK_CORRUPTION and encoding:
        try:
            content = preview if isinstance(preview, bytes) else _read_file_preview(path, 65536)
            text = content.decode(encoding or "utf-8", errors="replace")
            replacement_count = text.count("\ufffd")
            null_count = content.count(b"\x00")
            total_chars = len(text)
            if total_chars > 0:
                suspicious_ratio = replacement_count / total_chars
                if suspicious_ratio > SUSPICIOUS_CHARS_THRESHOLD:
                    result["status"] = "error"
                    result["issues"].append({
                        "severity": "error",
                        "type": "corrupt_encoding",
                        "message": f"High ratio of replacement characters: {suspicious_ratio:.2%} ({replacement_count}/{total_chars})"
                    })
                elif suspicious_ratio > 0.02:
                    result["issues"].append({
                        "severity": "warning",
                        "type": "suspicious_chars",
                        "message": f"Some suspicious characters found: {suspicious_ratio:.2%} ({replacement_count}/{total_chars})"
                    })
            if null_count > 0:
                result["issues"].append({
                    "severity": "warning",
                    "type": "null_bytes",
                    "message": f"File contains {null_count} null byte(s)"
                })
        except Exception as e:
            result["issues"].append({
                "severity": "warning",
                "type": "content_check_error",
                "message": f"Error reading text content: {str(e)}"
            })


def _check_pdf(path, result):
    if not CHECK_OPENABLE and not CHECK_CORRUPTION:
        return
    try:
        import fitz
        doc = fitz.open(path)
        result["page_count"] = doc.page_count
        has_text = False
        for page_num in range(min(doc.page_count, 5)):
            try:
                text = doc[page_num].get_text()
                if text and text.strip():
                    has_text = True
                    break
            except Exception:
                continue
        if doc.page_count > 0 and not has_text:
            result["issues"].append({
                "severity": "warning",
                "type": "scanned_pdf",
                "message": f"PDF appears to be scanned (no extractable text in first 5 pages of {doc.page_count})"
            })
        doc.close()
    except ImportError:
        logger.warning("PyMuPDF not available, skipping PDF check for %s", path)
        result["issues"].append({
            "severity": "warning",
            "type": "library_missing",
            "message": "PyMuPDF (fitz) not installed, cannot check PDF"
        })
    except Exception as e:
        result["status"] = "error"
        result["issues"].append({
            "severity": "error",
            "type": "corrupt_pdf",
            "message": f"Cannot open PDF: {str(e)}"
        })


def _check_image(path, result):
    if not CHECK_OPENABLE and not CHECK_CORRUPTION:
        return
    try:
        from PIL import Image
        img = Image.open(path)
        result["image_dimensions"] = (img.width, img.height)
        try:
            img.load()
        except Exception:
            result["issues"].append({
                "severity": "error",
                "type": "truncated_image",
                "message": f"Image appears truncated or corrupt (cannot fully load pixels)"
            })
            return
        result["page_count"] = getattr(img, "n_frames", 1)
    except ImportError:
        logger.warning("Pillow not available, skipping image check for %s", path)
        result["issues"].append({
            "severity": "warning",
            "type": "library_missing",
            "message": "Pillow not installed, cannot check image"
        })
    except Exception as e:
        result["status"] = "error"
        result["issues"].append({
            "severity": "error",
            "type": "corrupt_image",
            "message": f"Cannot open image: {str(e)}"
        })


def _check_office(path, result, category):
    if not CHECK_OPENABLE and not CHECK_CORRUPTION:
        return
    ext = result["ext"].lower()
    checked = False

    if ext == ".docx":
        try:
            import docx
            doc = docx.Document(path)
            result["page_count"] = len(doc.paragraphs)
            checked = True
        except ImportError:
            pass
        except Exception as e:
            result["issues"].append({
                "severity": "error",
                "type": "corrupt_docx",
                "message": f"Cannot open DOCX: {str(e)}"
            })
            checked = True

    if ext == ".xlsx" and not checked:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            result["page_count"] = len(wb.sheetnames)
            wb.close()
            checked = True
        except ImportError:
            pass
        except Exception as e:
            result["issues"].append({
                "severity": "error",
                "type": "corrupt_xlsx",
                "message": f"Cannot open XLSX: {str(e)}"
            })
            checked = True

    if ext == ".rtf" and not checked:
        try:
            from striprtf.striprtf import rtf_to_text
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                text = rtf_to_text(f.read())
            if text:
                result["page_count"] = max(1, len(text) // 2000)
            checked = True
        except ImportError:
            pass
        except Exception as e:
            result["issues"].append({
                "severity": "warning",
                "type": "corrupt_rtf",
                "message": f"Cannot parse RTF: {str(e)}"
            })
            checked = True

    if ext == ".doc" and not checked:
        try:
            subprocess.run(
                ["antiword", path],
                capture_output=True, timeout=10
            )
            checked = True
        except FileNotFoundError:
            pass
        except subprocess.TimeoutExpired:
            result["issues"].append({
                "severity": "warning",
                "type": "antiword_timeout",
                "message": "antiword timed out on .doc file"
            })
            checked = True
        except Exception as e:
            result["issues"].append({
                "severity": "warning",
                "type": "antiword_error",
                "message": f"antiword failed: {str(e)}"
            })
            checked = True

    if not checked:
        try:
            libre_paths = ["libreoffice", "soffice"]
            libre_cmd = None
            for cmd in libre_paths:
                try:
                    subprocess.run([cmd, "--version"], capture_output=True, timeout=5)
                    libre_cmd = cmd
                    break
                except (FileNotFoundError, subprocess.TimeoutExpired):
                    continue
            if libre_cmd:
                subprocess.run(
                    [libre_cmd, "--headless", "--norestore", "--nofirststartwizard",
                     "--convert-to", "txt", "--outdir", "/tmp", path],
                    capture_output=True, timeout=30
                )
                checked = True
        except FileNotFoundError:
            pass
        except subprocess.TimeoutExpired:
            pass
        except Exception as e:
            logger.debug("LibreOffice fallback failed: %s", e)

    if not checked:
        result["issues"].append({
            "severity": "warning",
            "type": "office_not_checked",
            "message": f"No checker available for {ext} files"
        })


def _check_audio_video(path, result):
    if not CHECK_OPENABLE and not CHECK_CORRUPTION:
        return
    try:
        proc = subprocess.run(
            ["ffprobe", "-v", "quiet", "-print_format", "json", "-show_format", "-show_streams", path],
            capture_output=True, timeout=60
        )
        if proc.returncode != 0:
            result["issues"].append({
                "severity": "error",
                "type": "corrupt_media",
                "message": f"ffprobe could not read file (may be corrupt)"
            })
            return
        import json
        data = json.loads(proc.stdout)
        if "format" in data:
            fmt = data["format"]
            duration = fmt.get("duration")
            if duration:
                result["page_count"] = int(float(duration))
            if fmt.get("nb_streams"):
                result["image_dimensions"] = (fmt["nb_streams"], 0)
        if "streams" in data:
            for stream in data["streams"]:
                codec = stream.get("codec_name", "unknown")
                if stream.get("codec_type") == "video":
                    w = stream.get("width", 0)
                    h = stream.get("height", 0)
                    if w and h:
                        result["image_dimensions"] = (w, h)
    except FileNotFoundError:
        result["issues"].append({
            "severity": "warning",
            "type": "ffprobe_missing",
            "message": "ffprobe not installed, cannot check audio/video"
        })
    except subprocess.TimeoutExpired:
        result["issues"].append({
            "severity": "warning",
            "type": "ffprobe_timeout",
            "message": "ffprobe timed out on media file"
        })
    except json.JSONDecodeError:
        result["issues"].append({
            "severity": "warning",
            "type": "ffprobe_parse_error",
            "message": "Could not parse ffprobe output"
        })
    except Exception as e:
        result["issues"].append({
            "severity": "warning",
            "type": "media_check_error",
            "message": f"Error checking media file: {str(e)}"
        })


def _check_archive(path, result):
    if not CHECK_OPENABLE and not CHECK_CORRUPTION:
        return
    ext = result["ext"].lower()
    checked = False

    if ext == ".zip":
        try:
            import zipfile
            with zipfile.ZipFile(path, "r") as zf:
                bad = zf.testzip()
                if bad:
                    result["issues"].append({
                        "severity": "error",
                        "type": "corrupt_zip",
                        "message": f"ZIP file is corrupt: {bad}"
                    })
                else:
                    result["page_count"] = len(zf.namelist())
            checked = True
        except zipfile.BadZipFile as e:
            result["issues"].append({
                "severity": "error",
                "type": "corrupt_zip",
                "message": f"Bad ZIP file: {str(e)}"
            })
            checked = True
        except Exception as e:
            result["issues"].append({
                "severity": "error",
                "type": "zip_error",
                "message": f"Error reading ZIP: {str(e)}"
            })
            checked = True

    if ext in (".tar", ".gz", ".bz2", ".xz") and not checked:
        try:
            import tarfile
            with tarfile.open(path, "r") as tf:
                members = tf.getmembers()
                result["page_count"] = len(members)
            checked = True
        except tarfile.ReadError as e:
            result["issues"].append({
                "severity": "error",
                "type": "corrupt_tar",
                "message": f"Bad TAR file: {str(e)}"
            })
            checked = True
        except Exception as e:
            result["issues"].append({
                "severity": "error",
                "type": "tar_error",
                "message": f"Error reading TAR: {str(e)}"
            })
            checked = True

    if not checked:
        result["issues"].append({
            "severity": "warning",
            "type": "archive_not_checked",
            "message": f"No checker available for {ext} archive"
        })


def _check_ebook(path, result):
    if not CHECK_OPENABLE and not CHECK_CORRUPTION:
        return
    ext = result["ext"].lower()
    try:
        if ext == ".epub":
            try:
                import zipfile
                with zipfile.ZipFile(path, "r") as zf:
                    names = zf.namelist()
                    if any(n.endswith(".opf") for n in names):
                        result["page_count"] = len(names)
                    else:
                        result["issues"].append({
                            "severity": "warning",
                            "type": "invalid_epub",
                            "message": "EPUB file missing .opf metadata"
                        })
            except zipfile.BadZipFile:
                result["issues"].append({
                    "severity": "error",
                    "type": "corrupt_epub",
                    "message": "EPUB file is corrupt (not a valid ZIP)"
                })
        elif ext == ".mobi":
            preview = _read_file_preview(path, 100)
            if len(preview) < 68:
                result["issues"].append({
                    "severity": "error",
                    "type": "corrupt_mobi",
                    "message": "MOBI file too small or corrupt"
                })
            else:
                magic = struct.unpack(">4s", preview[:4])[0]
                if magic != b"BOOK":
                    result["issues"].append({
                        "severity": "warning",
                        "type": "invalid_mobi",
                        "message": f"Invalid MOBI header magic: {magic}"
                    })
        elif ext == ".djvu":
            preview = _read_file_preview(path, 20)
            if not preview.startswith(b"AT&T"):
                result["issues"].append({
                    "severity": "warning",
                    "type": "invalid_djvu",
                    "message": "Invalid DJVU header magic"
                })
        elif ext in (".cbr", ".cbz"):
            try:
                import zipfile
                with zipfile.ZipFile(path, "r") as zf:
                    result["page_count"] = len(zf.namelist())
            except zipfile.BadZipFile:
                try:
                    import tarfile
                    with tarfile.open(path, "r") as tf:
                        result["page_count"] = len(tf.getmembers())
                except Exception:
                    result["issues"].append({
                        "severity": "error",
                        "type": "corrupt_comic",
                        "message": "Comic archive is corrupt or unsupported format"
                    })
        else:
            result["issues"].append({
                "severity": "warning",
                "type": "ebook_not_checked",
                "message": f"No checker available for {ext} ebook format"
            })
    except Exception as e:
        result["issues"].append({
            "severity": "error",
            "type": "ebook_error",
            "message": f"Error checking ebook: {str(e)}"
        })


def _check_binary(path, result):
    MAGIC_MAP = {
        b"\x7fELF": "ELF",
        b"MZ": "PE (DOS/Windows executable)",
        b"\x89PNG": "PNG image",
        b"\xff\xd8\xff": "JPEG image",
        b"GIF8": "GIF image",
        b"%PDF": "PDF document",
        b"\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00": None,
    }
    try:
        preview = _read_file_preview(path, 16)
        detected = None
        for magic, name in MAGIC_MAP.items():
            if preview.startswith(magic):
                detected = name
                break
        if detected:
            result["image_dimensions"] = (len(preview), 0)
        else:
            result["issues"].append({
                "severity": "info",
                "type": "unknown_binary",
                "message": f"Unknown binary file type, size: {result['size']} bytes"
            })
        if result["size"] == 0:
            result["issues"].append({
                "severity": "warning",
                "type": "empty_file",
                "message": "Binary file is empty"
            })
    except Exception as e:
        result["issues"].append({
            "severity": "warning",
            "type": "binary_check_error",
            "message": f"Error checking binary: {str(e)}"
        })
